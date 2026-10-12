#!/usr/bin/env python3
"""Mostra a estrutura de um .xlsx para entender o modelo antes de reproduzi-lo.

Imprime, para cada aba: dimensões, configuração de impressão, quebras de página,
larguras de coluna (inclusive grupos e colunas ocultas), alturas de linha,
células mescladas e o conteúdo das células. Com --estilos, também fonte,
preenchimento, alinhamento e formato numérico de cada célula com conteúdo.

Uso:
    python inspecionar_xlsx.py arquivo.xlsx [--aba NOME] [--linhas 1-40]
                               [--estilos] [--max-texto 300]

Dica: rode primeiro sem --linhas para ver o tamanho; depois foque na primeira
página (onde costuma estar o bloco que se repete) com --linhas e --estilos.
"""

import argparse
import re
import zipfile

from openpyxl import load_workbook


def resumo_zip(caminho):
    with zipfile.ZipFile(caminho) as z:
        nomes = z.namelist()
        print("== Pacote")
        print("imagens:", len([n for n in nomes if n.startswith("xl/media/")]),
              "| desenhos:", len([n for n in nomes if n.startswith("xl/drawings/drawing")]),
              "| links externos:", len([n for n in nomes if n.startswith("xl/externalLinks/") and n.endswith(".xml")]))
        for n in nomes:
            if re.match(r"xl/drawings/drawing\d+\.xml$", n):
                x = z.read(n).decode("utf-8", "ignore")
                print(f"{n}: fotos={x.count('<xdr:pic>')} formas={x.count('<xdr:sp ') + x.count('<xdr:sp>')} "
                      f"caixas de texto vazias≈{len(re.findall(r'txBox=\"1\"', x)) - x.count('<a:t>')}")


def faixa(texto):
    if not texto:
        return None
    a, _, b = texto.partition("-")
    return int(a), int(b or a)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("arquivo")
    ap.add_argument("--aba")
    ap.add_argument("--linhas", help="ex.: 1-40")
    ap.add_argument("--estilos", action="store_true")
    ap.add_argument("--max-texto", type=int, default=300)
    args = ap.parse_args()

    resumo_zip(args.arquivo)
    wb = load_workbook(args.arquivo)
    print("nomes definidos:", list(wb.defined_names.keys()))
    lim = faixa(args.linhas)

    for ws in wb.worksheets:
        if args.aba and ws.title != args.aba:
            continue
        print(f"\n== Aba {ws.title!r}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
        ps = ws.page_setup
        print(f"impressão: área={ws.print_area} orientação={ps.orientation} papel={ps.paperSize} "
              f"escala={ps.scale} fitToWidth={ps.fitToWidth} fitToHeight={ps.fitToHeight} "
              f"fitToPage={ws.sheet_properties.pageSetUpPr.fitToPage if ws.sheet_properties.pageSetUpPr else None}")
        m = ws.page_margins
        print(f"margens: esq={m.left} dir={m.right} sup={m.top} inf={m.bottom} | "
              f"rodapé={ws.oddFooter.center.text!r} | títulos repetidos={ws.print_title_rows}")
        print("quebras de linha:", [b.id for b in ws.row_breaks.brk])
        print("colunas:")
        for cd in sorted(ws.column_dimensions.values(), key=lambda c: c.min or 0):
            if cd.min and cd.min <= 60:
                print(f"  {cd.min}-{cd.max}: largura={cd.width} oculta={cd.hidden}")
        r0, r1 = lim or (1, ws.max_row)
        print("alturas:", {r: ws.row_dimensions[r].height for r in range(r0, r1 + 1) if ws.row_dimensions[r].height})
        print("mesclas:")
        for mr in sorted(ws.merged_cells.ranges, key=lambda x: (x.min_row, x.min_col)):
            if r0 <= mr.min_row <= r1:
                print("  ", mr.coord)
        print("células:")
        for linha in ws.iter_rows(min_row=r0, max_row=r1):
            for c in linha:
                if c.value is None:
                    continue
                v = str(c.value).replace("\n", " ⏎ ")
                if len(v) > args.max_texto:
                    v = v[:args.max_texto] + "…"
                print(f"  {c.coordinate}: {v}")
                if args.estilos:
                    f, al = c.font, c.alignment
                    cor = f.color.rgb if f.color is not None and isinstance(f.color.rgb, str) else None
                    fg = c.fill.fgColor
                    preench = fg.rgb if isinstance(fg.rgb, str) else f"tema {fg.theme}"
                    print(f"      fonte={f.name} {f.sz} {'negrito' if f.b else ''} cor={cor} | "
                          f"preench={c.fill.fill_type}:{preench} | alinh={al.horizontal}/{al.vertical} "
                          f"quebra={al.wrap_text} rot={al.textRotation} | formato={c.number_format}")


if __name__ == "__main__":
    main()
