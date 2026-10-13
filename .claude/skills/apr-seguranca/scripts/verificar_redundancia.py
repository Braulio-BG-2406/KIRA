#!/usr/bin/env python3
"""Aponta linhas de medidas repetidas ou quase iguais numa APR gerada.

Uso:
    python verificar_redundancia.py "Nome da APR.xlsx" [--limite 0.80]

Compara todas as linhas da coluna "Medidas de precauções" de todas as etapas.
Repetições esperadas (não precisam de correção): a marcação "Acompanhamento
Full time responsável Ambev", que o modelo coloca em cada etapa com
acompanhamento. O resto deve ser corrigido no JSON (ou na biblioteca) antes de
entregar.
"""

import argparse
import itertools
import re
import unicodedata
from difflib import SequenceMatcher

from openpyxl import load_workbook

ESPERADAS = {"acompanhamento full time responsavel ambev"}


def normalizar(texto: str) -> str:
    t = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode().lower()
    t = re.sub(r"^[#\-•*\s]*(\d+(\.\d+)*\.?)?\s*", "", t)
    return " ".join(re.sub(r"[^a-z0-9 ]+", " ", t).split())


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--limite", type=float, default=0.80, help="semelhança mínima para apontar (0 a 1)")
    args = ap.parse_args()

    ws = load_workbook(args.xlsx).active
    linhas = []
    for r in range(1, ws.max_row + 1):
        etapa = ws.cell(r, 2).value
        if isinstance(etapa, int):
            for t in str(ws.cell(r, 9).value or "").split("\n"):
                n = normalizar(t)
                if len(n) > 15 and n not in ESPERADAS:
                    linhas.append((etapa, t.strip(), n))
    achados = 0
    for (e1, t1, n1), (e2, t2, n2) in itertools.combinations(linhas, 2):
        taxa = SequenceMatcher(None, n1, n2).ratio()
        if taxa >= args.limite:
            achados += 1
            print(f"[{taxa:.2f}] etapa {e1}: {t1[:120]}\n       etapa {e2}: {t2[:120]}")
    print(f"{achados} possível(is) redundância(s) com semelhança >= {args.limite:.2f}.")


if __name__ == "__main__":
    main()
