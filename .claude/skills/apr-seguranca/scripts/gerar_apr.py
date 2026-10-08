#!/usr/bin/env python3
"""Gera a planilha da APR (Análise Preliminar de Riscos) a partir de um JSON.

Uso:
    python gerar_apr.py dados.json [--saida PASTA] [--pdf] [--hoje AAAA-MM-DD]

As datas são calculadas aqui, não no JSON:
    - elaboração = hoje (fuso America/Sao_Paulo)
    - execução   = dia seguinte
    - validade   = 30/12 do ano da execução (ou do ano seguinte, se a
                   execução cair depois de 30/12)
"""

import argparse
import json
import re
import shutil
import subprocess
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

EMPRESA_PADRAO = "TJ MONTAGENS"

# Colunas da tabela de riscos: Nº | Etapa | Perigo/Risco | Possíveis danos | Medidas de controle
LARGURAS = {"A": 6, "B": 30, "C": 32, "D": 30, "E": 70}
ULTIMA_COL = "E"

COR_TITULO = "1F3864"
COR_SECAO = "D9E1F2"
COR_CABECALHO = "BDD7EE"

fino = Side(style="thin", color="000000")
BORDA = Border(left=fino, right=fino, top=fino, bottom=fino)
QUEBRA = Alignment(wrap_text=True, vertical="top")
CENTRO = Alignment(wrap_text=True, vertical="center", horizontal="center")


def calcular_datas(hoje: date) -> dict:
    execucao = hoje + timedelta(days=1)
    validade = date(execucao.year, 12, 30)
    if execucao > validade:
        validade = date(execucao.year + 1, 12, 30)
    fmt = "%d/%m/%Y"
    return {
        "elaboracao": hoje.strftime(fmt),
        "execucao": execucao.strftime(fmt),
        "validade": validade.strftime(fmt),
    }


def nome_arquivo(atividade: str) -> str:
    nome = re.sub(r'[<>:"/\\|?*\n\r\t]', " ", atividade)
    nome = re.sub(r"\s+", " ", nome).strip().rstrip(".")
    return (nome[:150] or "APR") + ".xlsx"


def estimar_altura(texto: str, largura_col: float) -> float:
    """Altura de linha aproximada para o texto caber com quebra automática."""
    chars_por_linha = max(int(largura_col * 1.2), 1)
    linhas = 0
    for trecho in str(texto).split("\n"):
        linhas += max(1, -(-len(trecho) // chars_por_linha))
    return max(15, linhas * 13.5 + 3)


class Planilha:
    def __init__(self, ws):
        self.ws = ws
        self.linha = 1

    def mesclar(self, col_ini, col_fim, valor, *, negrito=False, fill=None,
                alinhamento=QUEBRA, tamanho=10, cor_fonte="000000", altura=None):
        ws, r = self.ws, self.linha
        ws.merge_cells(f"{col_ini}{r}:{col_fim}{r}")
        c = ws[f"{col_ini}{r}"]
        c.value = valor
        c.font = Font(bold=negrito, size=tamanho, color=cor_fonte)
        c.alignment = alinhamento
        for col in range(ord(col_ini), ord(col_fim) + 1):
            cel = ws[f"{chr(col)}{r}"]
            cel.border = BORDA
            if fill:
                cel.fill = PatternFill("solid", fgColor=fill)
        if altura:
            ws.row_dimensions[r].height = altura
        return c

    def secao(self, titulo):
        self.mesclar("A", ULTIMA_COL, titulo, negrito=True, fill=COR_SECAO,
                     alinhamento=Alignment(vertical="center"), altura=18)
        self.linha += 1

    def campo(self, rotulo, valor, rotulo2=None, valor2=None):
        """Linha de cabeçalho: rótulo em A:B, valor em C (ou C:E se não houver par)."""
        r = self.linha
        self.mesclar("A", "B", rotulo, negrito=True, fill=COR_CABECALHO)
        if rotulo2 is None:
            self.mesclar("C", ULTIMA_COL, valor)
            altura = estimar_altura(valor, LARGURAS["C"] + LARGURAS["D"] + LARGURAS["E"])
        else:
            self.mesclar("C", "C", valor)
            self.mesclar("D", "D", rotulo2, negrito=True, fill=COR_CABECALHO)
            self.mesclar("E", "E", valor2)
            altura = max(estimar_altura(valor, LARGURAS["C"]),
                         estimar_altura(valor2, LARGURAS["E"]))
        self.ws.row_dimensions[r].height = max(18, altura)
        self.linha += 1

    def lista(self, itens, colunas=2):
        """Lista de itens com marcador, distribuída em colunas (A:C e D:E)."""
        itens = [f"☒ {i}" for i in itens if str(i).strip()]
        if not itens:
            itens = ["—"]
        metade = -(-len(itens) // colunas)
        blocos = [itens[i:i + metade] for i in range(0, len(itens), metade)]
        while len(blocos) < 2:
            blocos.append([])
        esq, dir_ = "\n".join(blocos[0]), "\n".join(blocos[1])
        self.mesclar("A", "C", esq)
        self.mesclar("D", ULTIMA_COL, dir_)
        altura = max(estimar_altura(esq, 68), estimar_altura(dir_, 100))
        self.ws.row_dimensions[self.linha].height = altura
        self.linha += 1

    def texto(self, valor):
        self.mesclar("A", ULTIMA_COL, valor)
        self.ws.row_dimensions[self.linha].height = estimar_altura(valor, 168)
        self.linha += 1

    def espaco(self):
        self.linha += 1


def montar(dados: dict, datas: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "APR"
    for col, larg in LARGURAS.items():
        ws.column_dimensions[col].width = larg

    p = Planilha(ws)

    p.mesclar("A", ULTIMA_COL, "ANÁLISE PRELIMINAR DE RISCO – APR", negrito=True,
              fill=COR_TITULO, cor_fonte="FFFFFF", tamanho=14, alinhamento=CENTRO, altura=28)
    p.linha += 1

    # 1. Identificação
    p.secao("1. IDENTIFICAÇÃO")
    p.campo("Empresa executante", dados.get("empresa") or EMPRESA_PADRAO,
            "Nº da APR", dados.get("numero", ""))
    p.campo("Contratante", dados.get("contratante", ""),
            "Área / Local", dados.get("local", ""))
    p.campo("Atividade", dados["atividade"])
    if dados.get("descricao"):
        p.campo("Descrição do serviço", dados["descricao"])
    p.campo("Data de elaboração", datas["elaboracao"],
            "Data de execução", datas["execucao"])
    p.campo("Validade do documento", f"Até {datas['validade']}",
            "Elaborado por", dados.get("responsavel", ""))
    p.espaco()

    # 2. Documentos / permissões
    p.secao("2. PERMISSÕES, LIBERAÇÕES E DOCUMENTOS NECESSÁRIOS")
    p.lista(dados.get("permissoes", []))
    p.espaco()

    # 3. Ferramentas e equipamentos
    p.secao("3. FERRAMENTAS, MÁQUINAS E EQUIPAMENTOS")
    p.lista(dados.get("ferramentas", []))
    p.espaco()

    # 4. Tabela de riscos
    p.secao("4. ANÁLISE DE RISCOS POR ETAPA DA ATIVIDADE")
    r = p.linha
    for col, titulo in zip("ABCDE", ["Nº", "Etapa da atividade", "Perigo / Risco",
                                      "Possíveis danos / consequências",
                                      "Medidas de controle"]):
        c = ws[f"{col}{r}"]
        c.value = titulo
        c.font = Font(bold=True, size=10)
        c.fill = PatternFill("solid", fgColor=COR_CABECALHO)
        c.alignment = CENTRO
        c.border = BORDA
    ws.row_dimensions[r].height = 30
    p.linha += 1

    for n, etapa in enumerate(dados["etapas"], start=1):
        riscos = etapa.get("riscos") or [{"perigo": "", "dano": "", "medidas": []}]
        ini = p.linha
        for risco in riscos:
            r = p.linha
            medidas = risco.get("medidas", [])
            if isinstance(medidas, str):
                medidas = [medidas]
            texto_medidas = "\n".join(f"• {m}" for m in medidas)
            ws[f"C{r}"] = risco.get("perigo", "")
            ws[f"D{r}"] = risco.get("dano", "")
            ws[f"E{r}"] = texto_medidas
            for col in "ABCDE":
                ws[f"{col}{r}"].border = BORDA
                ws[f"{col}{r}"].alignment = QUEBRA
                ws[f"{col}{r}"].font = Font(size=10)
            ws.row_dimensions[r].height = max(
                estimar_altura(risco.get("perigo", ""), LARGURAS["C"]),
                estimar_altura(risco.get("dano", ""), LARGURAS["D"]),
                estimar_altura(texto_medidas, LARGURAS["E"]),
            )
            p.linha += 1
        fim = p.linha - 1
        ws[f"A{ini}"] = n
        ws[f"B{ini}"] = etapa.get("etapa", "")
        ws[f"A{ini}"].alignment = CENTRO
        ws[f"B{ini}"].alignment = Alignment(wrap_text=True, vertical="center")
        ws[f"B{ini}"].font = Font(bold=True, size=10)
        if fim > ini:
            ws.merge_cells(f"A{ini}:A{fim}")
            ws.merge_cells(f"B{ini}:B{fim}")
    p.espaco()

    # 5. EPIs e EPCs
    p.secao("5. EPIs OBRIGATÓRIOS (com CA válido)")
    p.lista(dados.get("epis", []))
    p.espaco()
    p.secao("6. EPCs, SINALIZAÇÃO E ISOLAMENTO")
    p.lista(dados.get("epcs", []))
    p.espaco()

    # 7. Emergência
    p.secao("7. PROCEDIMENTOS EM CASO DE EMERGÊNCIA")
    p.texto("\n".join(f"• {e}" for e in dados.get("emergencia", [])) or "—")
    p.espaco()

    # 8. Recomendações gerais
    p.secao("8. RECOMENDAÇÕES GERAIS")
    p.texto("\n".join(f"• {e}" for e in dados.get("recomendacoes", [])) or "—")
    p.espaco()

    # 9. Equipe
    p.secao("9. EQUIPE EXECUTANTE – declaro que participei da análise, entendi os riscos e as medidas de controle")
    r = p.linha
    for col_ini, col_fim, titulo in [("A", "B", "Nome"), ("C", "C", "Função"),
                                     ("D", "D", "Matrícula / Documento"), ("E", "E", "Assinatura")]:
        p.mesclar(col_ini, col_fim, titulo, negrito=True, fill=COR_CABECALHO, alinhamento=CENTRO)
    p.linha += 1
    for _ in range(int(dados.get("linhas_equipe", 8))):
        for col_ini, col_fim in [("A", "B"), ("C", "C"), ("D", "D"), ("E", "E")]:
            p.mesclar(col_ini, col_fim, "")
        ws.row_dimensions[p.linha].height = 22
        p.linha += 1
    p.espaco()

    # 10. Aprovações
    p.secao("10. APROVAÇÕES")
    for col_ini, col_fim, titulo in [("A", "B", "Função"), ("C", "C", "Nome"),
                                     ("D", "D", "Data"), ("E", "E", "Assinatura")]:
        p.mesclar(col_ini, col_fim, titulo, negrito=True, fill=COR_CABECALHO, alinhamento=CENTRO)
    p.linha += 1
    for funcao in ["Elaborado por (Segurança do Trabalho)", "Encarregado / Supervisor da executante",
                   "Responsável da área (contratante)", "Segurança do Trabalho (contratante)"]:
        p.mesclar("A", "B", funcao, negrito=True)
        for col in "CDE":
            p.mesclar(col, col, "")
        ws.row_dimensions[p.linha].height = 24
        p.linha += 1

    # Impressão: A4 paisagem, ajustado à largura
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_area = f"A1:{ULTIMA_COL}{p.linha - 1}"
    ws.oddFooter.center.text = "APR – &[Page] de &N"
    ws.sheet_view.showGridLines = False
    return wb


def validar(dados: dict):
    erros = []
    if not str(dados.get("atividade", "")).strip():
        erros.append("campo 'atividade' vazio")
    etapas = dados.get("etapas") or []
    if not etapas:
        erros.append("nenhuma etapa em 'etapas'")
    for i, e in enumerate(etapas, 1):
        if not e.get("riscos"):
            erros.append(f"etapa {i} sem riscos")
        for j, r in enumerate(e.get("riscos", []), 1):
            if not r.get("perigo") or not r.get("medidas"):
                erros.append(f"etapa {i}, risco {j}: falta 'perigo' ou 'medidas'")
    if erros:
        sys.exit("JSON inválido:\n- " + "\n- ".join(erros))


def converter_pdf(xlsx: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("Aviso: LibreOffice não encontrado; PDF não gerado.", file=sys.stderr)
        return None
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir",
                    str(xlsx.parent), str(xlsx)], check=True, capture_output=True, timeout=180)
    pdf = xlsx.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("json", help="arquivo JSON com os dados da APR")
    ap.add_argument("--saida", default=".", help="pasta de saída (padrão: atual)")
    ap.add_argument("--pdf", action="store_true", help="também gera PDF via LibreOffice")
    ap.add_argument("--hoje", help="força a data de elaboração (AAAA-MM-DD); só para testes")
    args = ap.parse_args()

    dados = json.loads(Path(args.json).read_text(encoding="utf-8"))
    validar(dados)

    hoje = (date.fromisoformat(args.hoje) if args.hoje
            else datetime.now(ZoneInfo("America/Sao_Paulo")).date())
    datas = calcular_datas(hoje)

    saida = Path(args.saida)
    saida.mkdir(parents=True, exist_ok=True)
    destino = saida / nome_arquivo(dados["atividade"])
    montar(dados, datas).save(destino)
    print(f"XLSX: {destino}")
    print(f"Elaboração {datas['elaboracao']} | Execução {datas['execucao']} | Validade {datas['validade']}")
    if args.pdf:
        pdf = converter_pdf(destino)
        if pdf:
            print(f"PDF: {pdf}")


if __name__ == "__main__":
    main()
