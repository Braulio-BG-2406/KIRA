#!/usr/bin/env python3
"""EXEMPLO DE REFERÊNCIA (skill apr-seguranca) — não roda daqui: depende dos
assets daquela skill. Use como base para o gerador de um modelo novo.

Gera a APR no modelo "ANEXO 9 - ANÁLISE PRELIMINAR DE RISCOS / DECLARAÇÃO DE

Uso:
    python gerar_apr.py dados.json [--saida PASTA] [--pdf] [--hoje AAAA-MM-DD]

As datas são calculadas aqui, não no JSON:
    - DATA ELABORAÇÃO / EMISSÃO DA APR = hoje (fuso America/Sao_Paulo)
    - DATA PREVISTA DE INÍCIO          = dia seguinte
    - DATA PREVISTA DE TÉRMINO         = 30/12 do ano do início (ou do ano
                                         seguinte, se o início cair depois)

Cada página repete cabeçalho, títulos da tabela e bloco de assinaturas, como no
modelo. Etapas longas continuam na página seguinte com "Continuação...".
"""

import argparse
import json
import math
import re
import shutil
import subprocess
import sys
from copy import copy
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.pagebreak import Break
from PIL import Image as PILImage

RAIZ = Path(__file__).resolve().parent.parent
ASSETS = RAIZ / "assets"
FOTOS = ASSETS / "fotos"

# Linhas do molde (assets/modelo-anexo9.xlsx)
MOLDE_ESPACO = 1
MOLDE_CABECALHO = range(2, 9)
MOLDE_TITULOS = range(9, 13)
MOLDE_CONTEUDO = 13
MOLDE_ASSINATURAS = range(14, 19)
NCOLS = 19  # A..S

ALTURA_MAX_LINHA = 409       # limite do Excel para uma linha
CAPACIDADE_PAGINA = 585      # pontos de conteúdo por página, como no modelo original
ALTURA_TEXTO = 12.75         # altura de uma linha de Arial 10 no Excel
FOLGA = 12
# caracteres por linha (Arial 10) em cada coluna de conteúdo
LARGURA = {"seq": 24, "tecnicas": 54, "riscos": 26, "sif": 38, "medidas": 120}
LARGURA_FOTO_PX = 440        # coluna P (FOTOS ILUSTRATIVA)

SIF_SIM = "Precursor - Sim\nPotencial - Sim\nReal - Sim"
SIF_NAO = "Precursor - Não\nPotencial - Não\nReal - Não"


# ----------------------------------------------------------------- datas

def calcular_datas(hoje: date) -> dict:
    inicio = hoje + timedelta(days=1)
    termino = date(inicio.year, 12, 30)
    if inicio > termino:
        termino = date(inicio.year + 1, 12, 30)
    return {"elaboracao": hoje, "inicio": inicio, "termino": termino}


def nome_arquivo(atividade: str) -> str:
    nome = re.sub(r'[<>:"/\\|?*\n\r\t]', " ", atividade)
    nome = re.sub(r"\s+", " ", nome).strip().rstrip(". ")
    return (nome[:150] or "APR") + ".xlsx"


# ------------------------------------------------------- texto e altura

def linhas_visuais(texto: str, chars: int) -> int:
    total = 0
    for linha in str(texto).split("\n"):
        total += max(1, math.ceil(len(linha.rstrip()) * 1.05 / chars))
    return total


def fatiar(texto: str, chars: int, capacidade: int) -> list[str]:
    """Divide o texto em fatias que cabem numa linha da planilha, cortando de
    preferência entre parágrafos (blocos separados por linha em branco)."""
    grupos = [g.split("\n") for g in re.split(r"\n\s*\n", str(texto).strip())]
    fatias, atual, usado = [], [], 0
    for grupo in grupos:
        h = sum(linhas_visuais(l, chars) for l in grupo) + (1 if atual else 0)
        if usado + h <= capacidade:
            if atual:
                atual.append("")
            atual.extend(grupo)
            usado += h
            continue
        if atual:
            fatias.append(atual)
            atual, usado = [], 0
        for linha in grupo:
            hl = linhas_visuais(linha, chars)
            if atual and usado + hl > capacidade:
                fatias.append(atual)
                atual, usado = [], 0
            atual.append(linha)
            usado += hl
    if atual:
        fatias.append(atual)
    return ["\n".join(f).strip("\n") for f in fatias] or [""]


# ------------------------------------------------------------- conteúdo

def carregar_json(caminho: Path) -> dict:
    return json.loads(caminho.read_text(encoding="utf-8"))


def montar_atividade(ativ: dict, biblio: dict, cfg: dict, base_json: Path) -> dict:
    """Etapa da atividade principal a partir do catálogo de riscos."""
    catalogo = biblio["riscos"]
    escolhidos = []
    for chave in ativ.get("riscos", []):
        if chave not in catalogo:
            sys.exit(f"Risco desconhecido: '{chave}'. Válidos: {', '.join(catalogo)}")
        r = dict(catalogo[chave])
        r["medidas"] = list(r["medidas"]) + list(ativ.get("medidas_extras", {}).get(chave, []))
        escolhidos.append(r)
    for extra in ativ.get("riscos_extras", []):
        if extra.get("categoria") not in dict(biblio["categorias"]):
            sys.exit(f"Categoria inválida em riscos_extras: {extra.get('categoria')}")
        escolhidos.append(extra)

    linhas_riscos, blocos_medidas, fotos_por_titulo = [], [], {}
    n = 0
    for cat, rotulo in biblio["categorias"]:
        da_cat = [r for r in escolhidos if r["categoria"] == cat]
        if not da_cat:
            continue
        if linhas_riscos:
            linhas_riscos.append("")
        linhas_riscos.append(rotulo)
        for r in da_cat:
            n += 1
            linhas_riscos.append(f"{n}. {r['titulo']}")
            medidas = [f"{n}. {r['medidas'][0]}"] + [
                f"{n}.{i}. {m}" for i, m in enumerate(r["medidas"][1:], start=1)]
            blocos_medidas.append(r["titulo"] + "\n" + "\n".join(medidas))
            fotos_por_titulo[r["titulo"]] = r.get("fotos", [])
    if not escolhidos:
        sys.exit("A atividade precisa de pelo menos um risco em 'riscos' ou 'riscos_extras'.")

    passos = ativ.get("passos", [])
    tecnicas = "\n\n".join(f"{i} Passo: {p}" for i, p in enumerate(passos, start=1))
    ferramentas = ativ.get("ferramentas", [])
    if ferramentas:
        tecnicas += ("\n\n" if tecnicas else "") + "FERRAMENTAS:\n" + "\n".join(ferramentas)

    full_time = ativ.get("full_time", "SIM")
    cabeca = "Acompanhamento Full time responsável Ambev\n\n" if full_time == "SIM" else ""
    rodape = biblio["rodape_atividade"].format(**cfg)
    # o rodapé fica colado ao último risco para não sobrar sozinho numa página
    medidas = cabeca + "\n\n".join(blocos_medidas) + "\n" + rodape

    fotos_usuario = []
    for f in ativ.get("fotos", []):
        p = Path(f)
        if not p.is_absolute():
            p = (base_json.parent / p) if (base_json.parent / p).exists() else FOTOS / p
        if p.exists():
            fotos_usuario.append(str(p))
        else:
            print(f"Aviso: foto não encontrada: {f}", file=sys.stderr)

    return {
        "sequencia": ativ.get("sequencia") or ativ.get("atividade", ""),
        "tecnicas": tecnicas or "—",
        "riscos": "\n".join(linhas_riscos),
        "sif": ativ.get("sif", True),
        "medidas": medidas,
        "full_time": full_time,
        "pt": ativ.get("pt", "SIM"),
        "fotos": fotos_usuario,
        "fotos_por_titulo": fotos_por_titulo,
        "altura_min": 0,
    }


def montar_etapas(dados: dict, biblio: dict, cfg: dict, base_json: Path) -> list[dict]:
    padrao = biblio["etapas_padrao"]
    atividades = dados.get("atividades") or [dict(dados, sequencia=dados.get("sequencia") or dados["atividade"])]

    todas_ferramentas = []
    for a in atividades:
        for f in a.get("ferramentas", []):
            if f not in todas_ferramentas:
                todas_ferramentas.append(f)
    fmt = dict(cfg, ferramentas=", ".join(todas_ferramentas) or "ferramentas manuais")

    def de_padrao(chave):
        e = dict(padrao[chave])
        for campo in ("sequencia", "tecnicas", "riscos", "medidas"):
            e[campo] = e[campo].format(**fmt)
        e["fotos"] = [str(FOTOS / f) for f in e.get("fotos", []) if (FOTOS / f).exists()]
        return e

    etapas = [de_padrao("planejamento"), de_padrao("transporte_isolamento")]
    if dados.get("utilidades"):
        etapas.append(de_padrao("rota_fuga_nh3"))
    if dados.get("bloqueio_energia", True):
        etapas.append(de_padrao("bloqueio_energia"))
    for a in atividades:
        etapas.append(montar_atividade(a, biblio, cfg, base_json))
    for extra in dados.get("etapas_extras", []):
        etapas.append({
            "sequencia": extra["sequencia"], "tecnicas": extra.get("tecnicas", ""),
            "riscos": extra.get("riscos", ""), "sif": extra.get("sif", False),
            "medidas": extra.get("medidas", ""), "full_time": extra.get("full_time", "NÃO"),
            "pt": extra.get("pt", "NÃO"), "fotos": [], "altura_min": 0,
        })
    etapas.append(de_padrao("cinco_s"))
    for i, e in enumerate(etapas, start=1):
        e["num"] = i
    return etapas


def fatiar_etapa(etapa: dict) -> list[dict]:
    """Quebra uma etapa em blocos (linhas da planilha) que cabem numa página."""
    cap = int((CAPACIDADE_PAGINA - FOLGA) / ALTURA_TEXTO)
    col = {c: fatiar(etapa[c], LARGURA[c], cap) for c in ("tecnicas", "riscos", "medidas")}
    total = max(len(v) for v in col.values())
    sif = SIF_SIM if etapa["sif"] else SIF_NAO
    blocos = []
    for k in range(total):
        b = {
            "num": etapa["num"],
            "seq": etapa["sequencia"] if k == 0 else "Continuação...",
            "tecnicas": col["tecnicas"][k] if k < len(col["tecnicas"]) else "Continuação",
            "riscos": col["riscos"][k] if k < len(col["riscos"]) else "Continuação...",
            "sif": sif if k == 0 else "Continuação...",
            "medidas": col["medidas"][k] if k < len(col["medidas"]) else "Continuação...",
            "full_time": etapa["full_time"],
            "pt": etapa["pt"],
        }
        fotos = list(etapa.get("fotos", [])) if k == 0 else []
        for titulo, lista in etapa.get("fotos_por_titulo", {}).items():
            if titulo in b["medidas"].split("\n"):
                fotos += [str(FOTOS / f) for f in lista if (FOTOS / f).exists()]
        b["fotos"] = fotos[:4]
        linhas = max(linhas_visuais(b[c], LARGURA[c]) for c in ("seq", "tecnicas", "riscos", "sif", "medidas"))
        altura = linhas * ALTURA_TEXTO + FOLGA
        if k == 0:
            altura = max(altura, etapa.get("altura_min", 0))
        if b["fotos"]:
            altura = max(altura, 150 * math.ceil(len(b["fotos"]) / 2))
        b["altura"] = min(round(altura, 2), CAPACIDADE_PAGINA)
        blocos.append(b)
    return blocos


def paginar(blocos: list[dict]) -> list[list[dict]]:
    paginas, atual, usado = [], [], 0
    for b in blocos:
        if atual and usado + b["altura"] > CAPACIDADE_PAGINA:
            paginas.append(atual)
            atual, usado = [], 0
        atual.append(b)
        usado += b["altura"]
    if atual:
        paginas.append(atual)
    return paginas


# --------------------------------------------------------------- planilha

class Molde:
    """Guarda valores, estilos, alturas e mesclagens das linhas do molde."""

    def __init__(self, ws):
        self.celulas = {}
        for r in range(1, MOLDE_ASSINATURAS.stop):
            for c in range(1, NCOLS + 1):
                cel = ws.cell(r, c)
                self.celulas[(r, c)] = (cel.value, copy(cel._style))
        self.alturas = {r: ws.row_dimensions[r].height for r in range(1, MOLDE_ASSINATURAS.stop)}
        self.mesclas = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]

    def copiar_linha(self, ws, origem: int, destino: int, altura=None):
        for c in range(1, NCOLS + 1):
            valor, estilo = self.celulas[(origem, c)]
            cel = ws.cell(destino, c)
            cel._style = copy(estilo)
            cel.value = valor
        ws.row_dimensions[destino].height = altura if altura is not None else self.alturas[origem]

    def copiar_faixa(self, ws, linhas: range, destino: int) -> list:
        desloc = destino - linhas.start
        for r in linhas:
            self.copiar_linha(ws, r, r + desloc)
        return [(r1 + desloc, c1, r2 + desloc, c2) for (r1, c1, r2, c2) in self.mesclas
                if r1 in linhas and r2 in linhas]


def inserir_fotos(ws, fotos: list[str], linha: int, alturas_pt: list[float]):
    if not fotos:
        return
    alturas_px = [h * 96 / 72 for h in alturas_pt]
    total_px = sum(alturas_px)
    ncol = 1 if len(fotos) <= 2 else 2
    nlin = math.ceil(len(fotos) / ncol)
    cel_w = (LARGURA_FOTO_PX - 20) / ncol
    cel_h = (total_px - 20) / nlin
    for i, caminho in enumerate(fotos):
        with PILImage.open(caminho) as im:
            w0, h0 = im.size
        escala = min(cel_w / w0, cel_h / h0, 1.0)
        w, h = max(int(w0 * escala), 1), max(int(h0 * escala), 1)
        x = 10 + (i % ncol) * cel_w + (cel_w - w) / 2
        y = 10 + (i // ncol) * cel_h + (cel_h - h) / 2
        r, y_rel = linha, y
        for hpx in alturas_px[:-1]:
            if y_rel >= hpx:
                y_rel -= hpx
                r += 1
        img = XLImage(caminho)
        img.width, img.height = w, h
        img.anchor = OneCellAnchor(
            _from=AnchorMarker(col=15, colOff=pixels_to_EMU(int(x)), row=r - 1, rowOff=pixels_to_EMU(int(y_rel))),
            ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)),
        )
        ws.add_image(img)


def gerar_planilha(dados: dict, cfg: dict, paginas: list, datas: dict, destino: Path):
    wb = load_workbook(ASSETS / "modelo-anexo9.xlsx")
    ws = wb.worksheets[0]
    molde = Molde(ws)
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(m.coord)
    ws.row_breaks.brk = []

    mesclas, linha = [], 1
    for np_, pagina in enumerate(paginas):
        if np_ == 0:
            molde.copiar_linha(ws, MOLDE_ESPACO, linha)
            linha += 1
        # cabeçalho
        ini = linha
        mesclas += molde.copiar_faixa(ws, MOLDE_CABECALHO, ini)
        d = ini - MOLDE_CABECALHO.start  # deslocamento em relação ao molde
        for (r, c), valor in {(3, 9): datas["elaboracao"], (4, 9): datas["inicio"], (4, 16): datas["termino"]}.items():
            cel = ws.cell(r + d, c)
            cel.value = valor
            cel.number_format = "DD/MM/YYYY"
        ws.cell(3 + d, 14).value = cfg.get("revisao", "0")
        ws.cell(5 + d, 6).value = dados.get("descricao") or dados["atividade"]
        ws.cell(6 + d, 6).value = dados.get("responsavel") or cfg["responsavel"]
        ws.cell(6 + d, 13).value = dados.get("telefone_responsavel") or cfg["telefone_responsavel"]
        ws.cell(7 + d, 6).value = dados.get("responsavel_seguranca") or cfg["responsavel_seguranca"]
        ws.cell(7 + d, 13).value = dados.get("telefone_seguranca") or cfg["telefone_seguranca"]
        ws.cell(8 + d, 6).value = dados.get("local", "")
        linha += len(MOLDE_CABECALHO)
        mesclas += molde.copiar_faixa(ws, MOLDE_TITULOS, linha)
        linha += len(MOLDE_TITULOS)

        # blocos de conteúdo
        for b in pagina:
            duas = b["altura"] > ALTURA_MAX_LINHA
            alturas = [b["altura"] / 2, b["altura"] / 2] if duas else [b["altura"]]
            for k, h in enumerate(alturas):
                molde.copiar_linha(ws, MOLDE_CONTEUDO, linha + k, altura=h)
            ws.cell(linha, 2).value = b["num"]
            ws.cell(linha, 3).value = b["seq"]
            ws.cell(linha, 5).value = b["tecnicas"]
            ws.cell(linha, 7).value = b["riscos"]
            ws.cell(linha, 8).value = b["sif"]
            ws.cell(linha, 9).value = b["medidas"]
            ws.cell(linha, 17).value = b["full_time"]
            ws.cell(linha, 18).value = b["pt"]
            fim = linha + len(alturas) - 1
            for c1, c2 in [(2, 2), (3, 4), (5, 6), (7, 7), (8, 8), (9, 15), (16, 16), (17, 17), (18, 18)]:
                if c1 != c2 or fim > linha:
                    mesclas.append((linha, c1, fim, c2))
            inserir_fotos(ws, b["fotos"], linha, alturas)
            linha = fim + 1

        # assinaturas
        mesclas += molde.copiar_faixa(ws, MOLDE_ASSINATURAS, linha)
        linha += len(MOLDE_ASSINATURAS)
        ws.row_breaks.append(Break(id=linha - 1))

    for r1, c1, r2, c2 in mesclas:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    ws.print_area = f"A1:R{linha - 1}"
    wb.save(destino)


# ------------------------------------------------------------------- main

def validar(dados: dict):
    erros = []
    if not str(dados.get("atividade", "")).strip():
        erros.append("campo 'atividade' vazio")
    atividades = dados.get("atividades") or [dados]
    for i, a in enumerate(atividades, 1):
        if not a.get("passos"):
            erros.append(f"atividade {i}: faltam os 'passos'")
        if not (a.get("riscos") or a.get("riscos_extras")):
            erros.append(f"atividade {i}: faltam os 'riscos'")
    if erros:
        sys.exit("JSON inválido:\n- " + "\n- ".join(erros))


def converter_pdf(xlsx: Path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("Aviso: LibreOffice não encontrado; PDF não gerado.", file=sys.stderr)
        return None
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(xlsx.parent), str(xlsx)],
                   check=True, capture_output=True, timeout=300)
    pdf = xlsx.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("json", help="arquivo JSON com os dados da APR")
    ap.add_argument("--saida", default=".", help="pasta de saída (padrão: atual)")
    ap.add_argument("--pdf", action="store_true", help="também gera PDF via LibreOffice")
    ap.add_argument("--hoje", help="força a data de elaboração (AAAA-MM-DD); só para testes")
    args = ap.parse_args()

    caminho = Path(args.json).resolve()
    dados = carregar_json(caminho)
    validar(dados)
    biblio = carregar_json(ASSETS / "biblioteca.json")
    cfg = carregar_json(ASSETS / "config.json")

    hoje = date.fromisoformat(args.hoje) if args.hoje else datetime.now(ZoneInfo("America/Sao_Paulo")).date()
    datas = calcular_datas(hoje)

    blocos = [b for e in montar_etapas(dados, biblio, cfg, caminho) for b in fatiar_etapa(e)]
    paginas = paginar(blocos)

    saida = Path(args.saida)
    saida.mkdir(parents=True, exist_ok=True)
    destino = saida / nome_arquivo(dados["atividade"])
    gerar_planilha(dados, cfg, paginas, datas, destino)
    fmt = "%d/%m/%Y"
    print(f"XLSX: {destino} ({len(paginas)} páginas)")
    print(f"Elaboração {datas['elaboracao']:{fmt}} | Início {datas['inicio']:{fmt}} | Término {datas['termino']:{fmt}}")
    if args.pdf:
        pdf = converter_pdf(destino)
        if pdf:
            print(f"PDF: {pdf}")


if __name__ == "__main__":
    main()
